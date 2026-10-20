"""Guardado de todos los datos en planillas de Excel.

Reglas que protegen los datos:
- Nada se borra: las ventas se anulan (queda el registro) y los productos
  se dan de baja.
- Cada cambio se escribe primero en el registro (registro/AAAA-MM.jsonl,
  solo se le agregan líneas) y después en el Excel. Si se corta la luz o el
  Excel está abierto, al volver a abrir el programa se recupera todo.
- Los Excel se guardan "de golpe" (archivo temporal + reemplazo), así nunca
  queda una planilla a medio escribir.
- Las hojas están protegidas con contraseña: se pueden mirar pero no editar.
- Se hace un respaldo .zip por día en la carpeta respaldos/.
"""
import datetime as dt
import json
import os
import shutil
import threading
import unicodedata
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.protection import WorkbookProtection

CLAVE_HOJAS = "kiosco-pos-protegido"
RESPALDOS_A_GUARDAR = 60
DINERO = '"$" #,##0.00'
TEXTO = "@"

# Cada hoja: (clave interna, título de columna, ancho, formato)
HOJAS_GENERAL = {
    "Productos": [
        ("codigo", "Código", 18, TEXTO),
        ("nombre", "Producto", 36, None),
        ("categoria", "Categoría", 16, None),
        ("costo", "Costo", 12, DINERO),
        ("precio", "Precio", 12, DINERO),
        ("stock", "Stock", 10, None),
        ("stock_minimo", "Stock mínimo", 13, None),
        ("activo", "Activo", 8, None),
        ("actualizado", "Actualizado", 20, None),
    ],
}

HOJAS_MES = {
    "Ventas": [
        ("numero", "N° venta", 10, None),
        ("fecha", "Fecha", 20, None),
        ("caja", "N° caja", 9, None),
        ("total", "Total", 13, DINERO),
        ("medio_pago", "Medio de pago", 18, None),
        ("pago_con", "Pagó con", 12, DINERO),
        ("vuelto", "Vuelto", 12, DINERO),
        ("articulos", "Artículos", 10, None),
        ("estado", "Estado", 11, None),
        ("anulada_fecha", "Fecha anulación", 20, None),
        ("motivo", "Motivo anulación", 30, None),
    ],
    "Detalle": [
        ("numero", "N° venta", 10, None),
        ("fecha", "Fecha", 20, None),
        ("codigo", "Código", 18, TEXTO),
        ("nombre", "Producto", 36, None),
        ("cantidad", "Cantidad", 10, None),
        ("precio", "Precio", 12, DINERO),
        ("costo", "Costo", 12, DINERO),
        ("subtotal", "Subtotal", 13, DINERO),
    ],
    "Caja": [
        ("caja", "N° caja", 9, None),
        ("fecha", "Fecha", 20, None),
        ("tipo", "Movimiento", 12, None),
        ("medio", "Medio", 18, None),
        ("monto", "Monto", 13, DINERO),
        ("detalle", "Detalle", 50, None),
    ],
    "Stock": [
        ("fecha", "Fecha", 20, None),
        ("codigo", "Código", 18, TEXTO),
        ("nombre", "Producto", 36, None),
        ("cantidad", "Cantidad", 10, None),
        ("motivo", "Motivo", 28, None),
        ("stock_resultante", "Stock final", 12, None),
    ],
}


class ErrorDatos(Exception):
    """Error que se le puede mostrar tal cual al usuario."""


def ahora() -> str:
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def dinero(valor) -> float:
    return round(float(valor or 0), 2)


def cantidad(valor) -> float:
    numero = round(float(valor or 0), 3)
    return int(numero) if numero == int(numero) else numero


def normalizar(texto: str) -> str:
    """Minúsculas y sin acentos, para buscar."""
    texto = unicodedata.normalize("NFKD", str(texto or "").lower())
    return "".join(c for c in texto if not unicodedata.combining(c))


class _Archivo:
    """Una planilla de Excel cargada en memoria."""

    def __init__(self, ruta: Path, esquema: dict):
        self.ruta = ruta
        self.esquema = esquema
        self.hojas = {nombre: [] for nombre in esquema}
        self.meta = {}
        self.seq_guardado = 0  # último cambio que ya está en el Excel
        self.seq = 0  # último cambio aplicado en memoria
        self.sucio = False

    def cargar(self):
        if not self.ruta.exists():
            return
        libro = load_workbook(self.ruta, read_only=True, data_only=True)
        try:
            for nombre, columnas in self.esquema.items():
                if nombre not in libro.sheetnames:
                    continue
                filas = libro[nombre].iter_rows(values_only=True)
                encabezado = next(filas, None) or ()
                por_titulo = {titulo: clave for clave, titulo, _, _ in columnas}
                claves = [por_titulo.get(titulo) for titulo in encabezado]
                for fila in filas:
                    if not fila or all(v is None for v in fila):
                        continue
                    registro = {clave: None for clave, *_ in columnas}
                    for clave, valor in zip(claves, fila):
                        if clave:
                            registro[clave] = valor
                    self.hojas[nombre].append(registro)
            if "_sistema" in libro.sheetnames:
                for fila in libro["_sistema"].iter_rows(values_only=True):
                    if fila and fila[0]:
                        self.meta[fila[0]] = fila[1]
        finally:
            libro.close()
        self.seq_guardado = self.seq = int(self.meta.get("seq") or 0)

    def foto(self):
        hojas = {n: [dict(r) for r in filas] for n, filas in self.hojas.items()}
        meta = dict(self.meta, seq=self.seq)
        return hojas, meta, self.seq

    @staticmethod
    def escribir(ruta: Path, esquema: dict, hojas: dict, meta: dict, proteger=True):
        libro = Workbook()
        libro.remove(libro.active)
        negrita = Font(bold=True, color="FFFFFF")
        fondo = PatternFill("solid", fgColor="1F6F5C")
        for nombre, columnas in esquema.items():
            hoja = libro.create_sheet(nombre)
            hoja.append([titulo for _, titulo, _, _ in columnas])
            for celda in hoja[1]:
                celda.font = negrita
                celda.fill = fondo
            for registro in hojas.get(nombre, []):
                hoja.append([registro.get(clave) for clave, *_ in columnas])
            for i, (_, _, ancho, formato) in enumerate(columnas, start=1):
                letra = hoja.cell(row=1, column=i).column_letter
                hoja.column_dimensions[letra].width = ancho
                if formato:
                    for (celda,) in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                        celda.number_format = formato
            hoja.freeze_panes = "A2"
            if hoja.max_row > 1:
                hoja.auto_filter.ref = hoja.dimensions
            if proteger:
                hoja.protection.sheet = True
                hoja.protection.password = CLAVE_HOJAS
                hoja.protection.autoFilter = False
                hoja.protection.sort = False
        if meta:
            sistema = libro.create_sheet("_sistema")
            for clave, valor in meta.items():
                sistema.append([clave, valor])
            sistema.sheet_state = "hidden"
            if proteger:
                sistema.protection.sheet = True
                sistema.protection.password = CLAVE_HOJAS
        if proteger:
            libro.security = WorkbookProtection(workbookPassword=CLAVE_HOJAS, lockStructure=True)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        temporal = ruta.with_name(ruta.name + ".tmp")
        libro.save(temporal)
        os.replace(temporal, ruta)


class Almacen:
    def __init__(self, carpeta, guardado_automatico=True):
        self.carpeta = Path(carpeta)
        self.carpeta_registro = self.carpeta / "registro"
        self.carpeta_respaldos = self.carpeta / "respaldos"
        for c in (self.carpeta, self.carpeta_registro, self.carpeta_respaldos):
            c.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._meses = {}
        self.error_guardado = None
        self.carpeta_respaldo_extra = ""

        self.general = _Archivo(self.carpeta / "kiosco.xlsx", HOJAS_GENERAL)
        self.general.cargar()
        self._reindexar_productos()
        self.seq = max(self.general.seq, self._mes(self.mes_actual()).seq)
        self.ultima_fecha = ""
        self._recuperar_registro()

        self._hay_cambios = threading.Event()
        self._fin = threading.Event()
        self._hilo = None
        if guardado_automatico:
            self._hilo = threading.Thread(target=self._bucle_guardado, daemon=True)
            self._hilo.start()

    # ------------------------------------------------------------------
    # Archivos y registro
    # ------------------------------------------------------------------
    @staticmethod
    def mes_actual() -> str:
        return dt.date.today().strftime("%Y-%m")

    def ruta_mes(self, mes: str) -> Path:
        return self.carpeta / f"ventas_{mes.replace('-', '_')}.xlsx"

    def meses_disponibles(self):
        meses = {self.mes_actual()}
        for ruta in self.carpeta.glob("ventas_*.xlsx"):
            partes = ruta.stem.split("_")
            if len(partes) == 3:
                meses.add(f"{partes[1]}-{partes[2]}")
        return sorted(meses)

    def _mes(self, mes: str, seq=None):
        """Devuelve la planilla del mes. Con seq, solo si ese cambio falta."""
        archivo = self._meses.get(mes)
        if archivo is None:
            archivo = _Archivo(self.ruta_mes(mes), HOJAS_MES)
            archivo.cargar()
            self._meses[mes] = archivo
        return self._marcar(archivo, seq)

    def _gen(self, seq):
        return self._marcar(self.general, seq)

    @staticmethod
    def _marcar(archivo, seq):
        if seq is None:
            return archivo
        if seq <= archivo.seq_guardado:
            return None  # este cambio ya estaba guardado en el Excel
        archivo.seq = max(archivo.seq, seq)
        archivo.sucio = True
        return archivo

    def _reindexar_productos(self):
        self.productos = {}
        for p in self.general.hojas["Productos"]:
            p["codigo"] = str(p["codigo"])
            self.productos[p["codigo"]] = p

    def _recuperar_registro(self):
        """Vuelve a aplicar los cambios que no llegaron a guardarse en Excel."""
        archivos = sorted(self.carpeta_registro.glob("*.jsonl"))[-3:]
        for ruta in archivos:
            with open(ruta, encoding="utf-8") as f:
                for linea in f:
                    try:
                        evento = json.loads(linea)
                    except ValueError:
                        continue  # línea cortada por un corte de luz
                    self.seq = max(self.seq, evento["seq"])
                    self.ultima_fecha = max(self.ultima_fecha, evento["fecha"])
                    self._aplicar(evento)

    def _registrar(self, tipo: str, **datos):
        with self._lock:
            self.seq += 1
            evento = {"seq": self.seq, "fecha": datos.pop("fecha", None) or ahora(), "tipo": tipo, **datos}
            ruta = self.carpeta_registro / f"{evento['fecha'][:7]}.jsonl"
            with open(ruta, "a", encoding="utf-8") as f:
                f.write(json.dumps(evento, ensure_ascii=False) + "\n")
                f.flush()
                os.fsync(f.fileno())
            self.ultima_fecha = max(self.ultima_fecha, evento["fecha"])
            self._aplicar(evento)
        if self._hilo:
            self._hay_cambios.set()
        return evento

    def _aplicar(self, ev):
        seq, tipo = ev["seq"], ev["tipo"]
        general = self._gen(seq)
        if general:
            for clave, valor in ev.get("meta", {}).items():
                general.meta[clave] = valor
            for p in ev.get("productos", []):
                existente = self.productos.get(p["codigo"])
                if existente is None:
                    self.general.hojas["Productos"].append(dict(p))
                    self.productos[p["codigo"]] = self.general.hojas["Productos"][-1]
                else:
                    existente.update(p)
        filas = {n: ev[n.lower()] for n in ("Ventas", "Detalle", "Caja", "Stock") if ev.get(n.lower())}
        mes = self._mes(ev["fecha"][:7], seq) if filas else None
        if mes:
            for nombre, registros in filas.items():
                mes.hojas[nombre].extend(dict(r) for r in registros)
        if tipo == "anulacion":
            mes_venta = self._mes(ev["mes_venta"], seq)
            if mes_venta:
                for venta in mes_venta.hojas["Ventas"]:
                    if venta["numero"] == ev["numero"]:
                        venta.update(estado="ANULADA", anulada_fecha=ev["fecha"], motivo=ev["motivo"])

    # ------------------------------------------------------------------
    # Guardado en segundo plano
    # ------------------------------------------------------------------
    def _bucle_guardado(self):
        while not self._fin.is_set():
            self._hay_cambios.wait(timeout=10)
            self._hay_cambios.clear()
            self.guardar_pendientes()

    def guardar_pendientes(self) -> bool:
        with self._lock:
            pendientes = [(a, a.foto()) for a in (self.general, *self._meses.values()) if a.sucio]
        error = None
        for archivo, (hojas, meta, seq) in pendientes:
            try:
                _Archivo.escribir(archivo.ruta, archivo.esquema, hojas, meta)
            except PermissionError:
                error = (f"No se pudo guardar {archivo.ruta.name} porque está abierto en Excel. "
                         "Cerralo: las ventas están a salvo y se guardan solas al cerrarlo.")
                continue
            except OSError as e:
                error = f"No se pudo guardar {archivo.ruta.name}: {e}"
                continue
            with self._lock:
                archivo.seq_guardado = seq
                if archivo.seq == seq:
                    archivo.sucio = False
        self.error_guardado = error
        return error is None

    def cerrar(self):
        self._fin.set()
        self._hay_cambios.set()
        if self._hilo:
            self._hilo.join(timeout=5)
        self.guardar_pendientes()

    # ------------------------------------------------------------------
    # Productos
    # ------------------------------------------------------------------
    def lista_productos(self, texto="", incluir_bajas=False, solo_stock_bajo=False):
        buscado = normalizar(texto)
        with self._lock:
            resultado = []
            for p in self.productos.values():
                if not incluir_bajas and p.get("activo") == "NO":
                    continue
                if solo_stock_bajo and float(p.get("stock") or 0) > float(p.get("stock_minimo") or 0):
                    continue
                if buscado and buscado not in normalizar(f"{p['codigo']} {p['nombre']} {p.get('categoria') or ''}"):
                    continue
                resultado.append(dict(p))
        return sorted(resultado, key=lambda p: normalizar(p["nombre"]))

    def producto(self, codigo):
        with self._lock:
            p = self.productos.get(str(codigo).strip())
            return dict(p) if p else None

    def categorias(self):
        with self._lock:
            return sorted({p["categoria"] for p in self.productos.values() if p.get("categoria")})

    def guardar_producto(self, datos: dict, codigo_original=None):
        """Crea o modifica un producto. El stock se cambia con mover_stock()."""
        codigo = str(datos.get("codigo") or "").strip()
        nombre = str(datos.get("nombre") or "").strip()
        if not nombre:
            raise ErrorDatos("El producto necesita un nombre.")
        with self._lock:
            if codigo_original and codigo and codigo != codigo_original:
                raise ErrorDatos("El código de un producto no se puede cambiar. Creá uno nuevo y dale de baja al viejo.")
            codigo = codigo or codigo_original
            existente = self.productos.get(codigo) if codigo else None
            if existente and not codigo_original:
                raise ErrorDatos(f"Ya existe un producto con el código {codigo}: {existente['nombre']}.")
            meta = {}
            if not codigo:
                numero = int(self.general.meta.get("ultimo_interno") or 0) + 1
                codigo = f"INT{numero:05d}"
                meta["ultimo_interno"] = numero
            producto = {
                "codigo": codigo,
                "nombre": nombre,
                "categoria": str(datos.get("categoria") or "").strip(),
                "costo": dinero(datos.get("costo")),
                "precio": dinero(datos.get("precio")),
                "stock_minimo": cantidad(datos.get("stock_minimo")),
                "activo": "SI",
                "actualizado": ahora(),
            }
            if producto["precio"] < 0 or producto["costo"] < 0:
                raise ErrorDatos("Los precios no pueden ser negativos.")
            if existente is None:
                producto["stock"] = 0
            self._registrar("producto", productos=[producto], meta=meta)
            stock_inicial = cantidad(datos.get("stock"))
            if existente is None and stock_inicial:
                self.mover_stock(codigo, stock_inicial, "Stock inicial")
            return codigo

    def cambiar_precios(self, porcentaje: float, categoria=""):
        """Aumenta (o baja) precios un porcentaje. Devuelve cuántos cambió."""
        factor = 1 + float(porcentaje) / 100
        with self._lock:
            cambios = []
            for p in self.productos.values():
                if p.get("activo") == "NO" or (categoria and p.get("categoria") != categoria):
                    continue
                cambios.append({"codigo": p["codigo"], "precio": dinero(float(p["precio"] or 0) * factor),
                                "actualizado": ahora()})
            if cambios:
                self._registrar("precios", productos=cambios)
        return len(cambios)

    def dar_de_baja(self, codigo, activo=False):
        with self._lock:
            if codigo not in self.productos:
                raise ErrorDatos("No existe ese producto.")
            self._registrar("producto", productos=[{"codigo": codigo, "activo": "SI" if activo else "NO",
                                                     "actualizado": ahora()}])

    def mover_stock(self, codigo, diferencia, motivo):
        """Suma (ingreso de mercadería) o resta (ajuste) stock."""
        with self._lock:
            p = self.productos.get(codigo)
            if p is None:
                raise ErrorDatos("No existe ese producto.")
            nuevo = cantidad(float(p.get("stock") or 0) + float(diferencia))
            fecha = ahora()
            self._registrar(
                "stock", fecha=fecha,
                productos=[{"codigo": codigo, "stock": nuevo, "actualizado": fecha}],
                stock=[{"fecha": fecha, "codigo": codigo, "nombre": p["nombre"], "cantidad": cantidad(diferencia),
                        "motivo": motivo, "stock_resultante": nuevo}],
            )
            return nuevo

    # ------------------------------------------------------------------
    # Ventas
    # ------------------------------------------------------------------
    def registrar_venta(self, items, medio_pago, pago_con=None):
        """items: [{"codigo","nombre","cantidad","precio"}]. Devuelve la venta."""
        if not items:
            raise ErrorDatos("No hay productos en la venta.")
        with self._lock:
            caja = self.caja_actual()
            if caja is None:
                raise ErrorDatos("Primero tenés que abrir la caja.")
            fecha = ahora()
            numero = int(self.general.meta.get("ultimo_numero") or 0) + 1
            detalle, stock, productos, stocks = [], [], [], {}
            total = 0.0
            for item in items:
                cant = cantidad(item["cantidad"])
                if cant <= 0:
                    raise ErrorDatos(f"Cantidad inválida para {item['nombre']}.")
                precio = dinero(item["precio"])
                subtotal = dinero(cant * precio)
                total += subtotal
                p = self.productos.get(str(item.get("codigo") or ""))
                detalle.append({"numero": numero, "fecha": fecha, "codigo": item.get("codigo") or "",
                                "nombre": item["nombre"], "cantidad": cant, "precio": precio,
                                "costo": dinero(p["costo"]) if p else 0, "subtotal": subtotal})
                if p:
                    actual = stocks.get(p["codigo"], float(p.get("stock") or 0))
                    stocks[p["codigo"]] = cantidad(actual - cant)
                    stock.append({"fecha": fecha, "codigo": p["codigo"], "nombre": p["nombre"], "cantidad": -cant,
                                  "motivo": f"Venta N° {numero}", "stock_resultante": stocks[p["codigo"]]})
            for codigo, nuevo in stocks.items():
                productos.append({"codigo": codigo, "stock": nuevo})
            total = dinero(total)
            pago = dinero(pago_con) if pago_con not in (None, "") else None
            if pago is not None and pago < total:
                raise ErrorDatos("El pago es menor que el total.")
            venta = {"numero": numero, "fecha": fecha, "caja": caja["numero"], "total": total,
                     "medio_pago": medio_pago, "pago_con": pago,
                     "vuelto": dinero(pago - total) if pago is not None else None,
                     "articulos": cantidad(sum(d["cantidad"] for d in detalle)), "estado": "OK",
                     "anulada_fecha": None, "motivo": None}
            self._registrar("venta", fecha=fecha, meta={"ultimo_numero": numero}, productos=productos,
                            ventas=[venta], detalle=detalle, stock=stock)
            return dict(venta, items=detalle)

    def anular_venta(self, numero, fecha_venta, motivo):
        motivo = (motivo or "").strip()
        if not motivo:
            raise ErrorDatos("Escribí el motivo de la anulación.")
        mes_venta = str(fecha_venta)[:7]
        with self._lock:
            venta = next((v for v in self._mes(mes_venta).hojas["Ventas"] if v["numero"] == numero), None)
            if venta is None:
                raise ErrorDatos("No se encontró la venta.")
            if venta["estado"] == "ANULADA":
                raise ErrorDatos("Esa venta ya estaba anulada.")
            fecha = ahora()
            stocks, stock = {}, []
            for d in self.detalle_venta(numero, fecha_venta):
                p = self.productos.get(str(d["codigo"]))
                if not p:
                    continue
                actual = stocks.get(p["codigo"], float(p.get("stock") or 0))
                stocks[p["codigo"]] = cantidad(actual + float(d["cantidad"]))
                stock.append({"fecha": fecha, "codigo": p["codigo"], "nombre": p["nombre"],
                              "cantidad": cantidad(d["cantidad"]), "motivo": f"Anulación venta N° {numero}",
                              "stock_resultante": stocks[p["codigo"]]})
            self._registrar("anulacion", fecha=fecha, numero=numero, mes_venta=mes_venta, motivo=motivo,
                            productos=[{"codigo": c, "stock": s} for c, s in stocks.items()], stock=stock)

    def ventas_entre(self, desde: str, hasta: str):
        """Ventas con fecha entre desde y hasta (texto AAAA-MM-DD, incluidas)."""
        resultado = []
        with self._lock:
            for mes in self.meses_disponibles():
                if desde[:7] <= mes <= hasta[:7]:
                    for v in self._mes(mes).hojas["Ventas"]:
                        if desde <= str(v["fecha"])[:10] <= hasta:
                            resultado.append(dict(v))
        return sorted(resultado, key=lambda v: v["numero"])

    def detalle_venta(self, numero, fecha_venta):
        with self._lock:
            return [dict(d) for d in self._mes(str(fecha_venta)[:7]).hojas["Detalle"] if d["numero"] == numero]

    def detalles_entre(self, desde, hasta):
        validas = {v["numero"] for v in self.ventas_entre(desde, hasta) if v["estado"] == "OK"}
        resultado = []
        with self._lock:
            for mes in self.meses_disponibles():
                if desde[:7] <= mes <= hasta[:7]:
                    resultado += [dict(d) for d in self._mes(mes).hojas["Detalle"] if d["numero"] in validas]
        return resultado

    def resumen(self, desde, hasta):
        ventas = [v for v in self.ventas_entre(desde, hasta) if v["estado"] == "OK"]
        por_medio = {}
        for v in ventas:
            por_medio[v["medio_pago"]] = dinero(por_medio.get(v["medio_pago"], 0) + v["total"])
        productos = {}
        ganancia = 0.0
        for d in self.detalles_entre(desde, hasta):
            fila = productos.setdefault(d["nombre"], {"nombre": d["nombre"], "cantidad": 0, "total": 0})
            fila["cantidad"] = cantidad(fila["cantidad"] + float(d["cantidad"]))
            fila["total"] = dinero(fila["total"] + float(d["subtotal"]))
            if d["costo"]:
                ganancia += float(d["subtotal"]) - float(d["costo"]) * float(d["cantidad"])
        total = dinero(sum(v["total"] for v in ventas))
        return {
            "total": total,
            "cantidad": len(ventas),
            "promedio": dinero(total / len(ventas)) if ventas else 0,
            "ganancia": dinero(ganancia),
            "por_medio": por_medio,
            "mas_vendidos": sorted(productos.values(), key=lambda f: -f["total"]),
        }

    # ------------------------------------------------------------------
    # Caja
    # ------------------------------------------------------------------
    def caja_actual(self):
        with self._lock:
            texto = self.general.meta.get("caja_actual")
            return json.loads(texto) if texto else None

    def abrir_caja(self, monto_inicial):
        with self._lock:
            if self.caja_actual():
                raise ErrorDatos("La caja ya está abierta.")
            numero = int(self.general.meta.get("ultima_caja") or 0) + 1
            fecha = ahora()
            caja = {"numero": numero, "apertura": fecha, "inicial": dinero(monto_inicial)}
            self._registrar("caja", fecha=fecha,
                            meta={"ultima_caja": numero, "caja_actual": json.dumps(caja)},
                            caja=[{"caja": numero, "fecha": fecha, "tipo": "APERTURA", "medio": "Efectivo",
                                   "monto": caja["inicial"], "detalle": "Monto inicial"}])
            return caja

    def movimiento_caja(self, tipo, monto, detalle):
        """tipo: INGRESO o EGRESO (retiros, pagos a proveedores, etc.)."""
        monto = dinero(monto)
        if monto <= 0:
            raise ErrorDatos("El monto tiene que ser mayor a cero.")
        with self._lock:
            caja = self.caja_actual()
            if not caja:
                raise ErrorDatos("La caja está cerrada.")
            self._registrar("caja", caja=[{"caja": caja["numero"], "fecha": ahora(), "tipo": tipo,
                                           "medio": "Efectivo", "monto": monto, "detalle": detalle}])

    def resumen_caja(self):
        with self._lock:
            caja = self.caja_actual()
            if not caja:
                return None
            hoy = dt.date.today().isoformat()
            ventas = [v for v in self.ventas_entre(caja["apertura"][:10], hoy)
                      if v["caja"] == caja["numero"] and v["estado"] == "OK"]
            movimientos = []
            for mes in self.meses_disponibles():
                if mes >= caja["apertura"][:7]:
                    movimientos += [m for m in self._mes(mes).hojas["Caja"] if m["caja"] == caja["numero"]]
        por_medio = {}
        for v in ventas:
            por_medio[v["medio_pago"]] = dinero(por_medio.get(v["medio_pago"], 0) + v["total"])
        ingresos = dinero(sum(m["monto"] for m in movimientos if m["tipo"] == "INGRESO"))
        egresos = dinero(sum(m["monto"] for m in movimientos if m["tipo"] == "EGRESO"))
        return {
            **caja,
            "ventas": len(ventas),
            "total_ventas": dinero(sum(v["total"] for v in ventas)),
            "por_medio": por_medio,
            "ingresos": ingresos,
            "egresos": egresos,
            "efectivo_esperado": dinero(caja["inicial"] + por_medio.get("Efectivo", 0) + ingresos - egresos),
        }

    def cerrar_caja(self, efectivo_contado):
        with self._lock:
            resumen = self.resumen_caja()
            if not resumen:
                raise ErrorDatos("La caja ya está cerrada.")
            contado = dinero(efectivo_contado)
            diferencia = dinero(contado - resumen["efectivo_esperado"])
            filas = [{"caja": resumen["numero"], "fecha": ahora(), "tipo": "VENTAS", "medio": medio, "monto": monto,
                      "detalle": "Total vendido en el turno"} for medio, monto in resumen["por_medio"].items()]
            filas.append({"caja": resumen["numero"], "fecha": ahora(), "tipo": "CIERRE", "medio": "Efectivo",
                          "monto": contado,
                          "detalle": f"Esperado $ {resumen['efectivo_esperado']:.2f} · Diferencia $ {diferencia:.2f} · "
                                     f"{resumen['ventas']} ventas por $ {resumen['total_ventas']:.2f}"})
            self._registrar("caja", meta={"caja_actual": ""}, caja=filas)
        self.respaldar()
        return dict(resumen, contado=contado, diferencia=diferencia)

    # ------------------------------------------------------------------
    # Respaldos y exportaciones
    # ------------------------------------------------------------------
    def respaldar(self):
        """Arma respaldos/respaldo_AAAA-MM-DD.zip con todo. Devuelve la ruta."""
        self.guardar_pendientes()
        destino = self.carpeta_respaldos / f"respaldo_{dt.date.today().isoformat()}.zip"
        temporal = destino.with_suffix(".tmp")
        with zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED) as z:
            for ruta in self.carpeta.glob("*.xlsx"):
                z.write(ruta, ruta.name)
            for ruta in self.carpeta_registro.glob("*.jsonl"):
                z.write(ruta, f"registro/{ruta.name}")
            if (self.carpeta / "configuracion.json").exists():
                z.write(self.carpeta / "configuracion.json", "configuracion.json")
        os.replace(temporal, destino)
        for viejo in sorted(self.carpeta_respaldos.glob("respaldo_*.zip"))[:-RESPALDOS_A_GUARDAR]:
            viejo.unlink()
        if self.carpeta_respaldo_extra:
            try:
                extra = Path(self.carpeta_respaldo_extra)
                extra.mkdir(parents=True, exist_ok=True)
                shutil.copy2(destino, extra / destino.name)
            except OSError:
                pass  # pendrive desconectado, etc.: el respaldo local ya está hecho
        return destino

    def respaldo_de_hoy_hecho(self):
        return (self.carpeta_respaldos / f"respaldo_{dt.date.today().isoformat()}.zip").exists()

    def exportar(self, desde, hasta):
        """Excel editable (sin protección) con las ventas del período."""
        destino = self.carpeta / "exportaciones" / f"ventas_{desde}_al_{hasta}.xlsx"
        ventas = self.ventas_entre(desde, hasta)
        numeros = {v["numero"] for v in ventas}
        with self._lock:
            hojas = {"Ventas": ventas, "Detalle": [], "Caja": [], "Stock": []}
            for mes in self.meses_disponibles():
                if desde[:7] <= mes <= hasta[:7]:
                    datos = self._mes(mes).hojas
                    hojas["Detalle"] += [d for d in datos["Detalle"] if d["numero"] in numeros]
                    for nombre in ("Caja", "Stock"):
                        hojas[nombre] += [r for r in datos[nombre] if desde <= str(r["fecha"])[:10] <= hasta]
            productos = [dict(p) for p in self.productos.values()]
        esquema = dict(HOJAS_MES, **HOJAS_GENERAL)
        _Archivo.escribir(destino, esquema, dict(hojas, Productos=productos), {}, proteger=False)
        return destino

import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from pos.almacen import Almacen, ErrorDatos  # noqa: E402

COCA = {"codigo": "7790895000997", "nombre": "Coca Cola 500", "categoria": "Bebidas", "costo": 800,
        "precio": 1200, "stock": 24, "stock_minimo": 6}


def item(codigo, nombre, cantidad, precio):
    return {"codigo": codigo, "nombre": nombre, "cantidad": cantidad, "precio": precio}


class TestAlmacen(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.carpeta = Path(self.tmp.name)
        self.a = Almacen(self.carpeta, guardado_automatico=False)

    def tearDown(self):
        self.tmp.cleanup()

    def reabrir(self, guardar=True):
        if guardar:
            self.a.guardar_pendientes()
        self.a = Almacen(self.carpeta, guardado_automatico=False)

    def test_venta_descuenta_stock_y_queda_en_excel(self):
        self.a.guardar_producto(COCA)
        self.a.abrir_caja(1000)
        venta = self.a.registrar_venta([item(COCA["codigo"], "Coca", 2, 1200)], "Efectivo", 3000)
        self.assertEqual(venta["total"], 2400)
        self.assertEqual(venta["vuelto"], 600)
        self.assertEqual(self.a.producto(COCA["codigo"])["stock"], 22)
        self.a.guardar_pendientes()

        libro = load_workbook(self.a.ruta_mes(self.a.mes_actual()))
        filas = list(libro["Ventas"].iter_rows(values_only=True))
        self.assertEqual(filas[0][0], "N° venta")
        self.assertEqual(filas[1][3], 2400)
        self.assertTrue(libro["Ventas"].protection.sheet)
        productos = load_workbook(self.carpeta / "kiosco.xlsx")["Productos"]
        self.assertEqual(productos["A2"].value, COCA["codigo"])  # el código queda como texto

    def test_recupera_ventas_que_no_llegaron_al_excel(self):
        self.a.guardar_producto(COCA)
        self.a.abrir_caja(0)
        self.a.guardar_pendientes()
        self.a.registrar_venta([item(COCA["codigo"], "Coca", 1, 1200)], "Débito")
        self.reabrir(guardar=False)  # como si se cortara la luz antes de guardar
        self.assertEqual(len(self.a.ventas_entre("2000-01-01", "2999-12-31")), 1)
        self.assertEqual(self.a.producto(COCA["codigo"])["stock"], 23)
        self.reabrir()
        self.reabrir()  # reaplicar el registro no duplica nada
        self.assertEqual(len(self.a.ventas_entre("2000-01-01", "2999-12-31")), 1)
        self.assertEqual(self.a.producto(COCA["codigo"])["stock"], 23)

    def test_si_el_excel_esta_abierto_no_se_pierde_nada(self):
        self.a.guardar_producto(COCA)
        self.a.guardar_pendientes()
        original = Path.replace

        def bloqueado(*args, **kwargs):
            raise PermissionError("abierto en Excel")

        import pos.almacen as modulo
        modulo.os.replace, guardado = bloqueado, modulo.os.replace
        try:
            self.a.mover_stock(COCA["codigo"], 10, "Ingreso")
            self.assertFalse(self.a.guardar_pendientes())
            self.assertIn("abierto en Excel", self.a.error_guardado)
        finally:
            modulo.os.replace = guardado
        self.assertIs(Path.replace, original)
        self.assertTrue(self.a.guardar_pendientes())
        self.reabrir()
        self.assertEqual(self.a.producto(COCA["codigo"])["stock"], 34)

    def test_anular_devuelve_stock_y_no_borra(self):
        self.a.guardar_producto(COCA)
        self.a.abrir_caja(0)
        venta = self.a.registrar_venta([item(COCA["codigo"], "Coca", 3, 1200)], "Efectivo")
        self.a.anular_venta(venta["numero"], venta["fecha"], "Se equivocó de producto")
        with self.assertRaises(ErrorDatos):
            self.a.anular_venta(venta["numero"], venta["fecha"], "otra vez")
        self.reabrir()
        ventas = self.a.ventas_entre("2000-01-01", "2999-12-31")
        self.assertEqual(ventas[0]["estado"], "ANULADA")
        self.assertEqual(self.a.producto(COCA["codigo"])["stock"], 24)
        self.assertEqual(self.a.resumen("2000-01-01", "2999-12-31")["total"], 0)

    def test_caja(self):
        with self.assertRaises(ErrorDatos):
            self.a.registrar_venta([item("", "Varios", 1, 500)], "Efectivo")
        self.a.abrir_caja(2000)
        self.a.registrar_venta([item("", "Varios", 1, 500)], "Efectivo")
        self.a.registrar_venta([item("", "Varios", 1, 700)], "Mercado Pago / QR")
        self.a.movimiento_caja("EGRESO", 300, "Pan")
        resumen = self.a.resumen_caja()
        self.assertEqual(resumen["efectivo_esperado"], 2200)
        cierre = self.a.cerrar_caja(2100)
        self.assertEqual(cierre["diferencia"], -100)
        self.assertIsNone(self.a.caja_actual())
        self.assertTrue(self.a.respaldo_de_hoy_hecho())
        self.a.abrir_caja(0)
        self.assertEqual(self.a.caja_actual()["numero"], 2)

    def test_productos(self):
        codigo = self.a.guardar_producto({"nombre": "Pan casero", "precio": 100, "categoria": "Panadería"})
        self.assertEqual(codigo, "INT00001")
        with self.assertRaises(ErrorDatos):
            self.a.guardar_producto(dict(COCA, codigo=codigo))
        self.a.guardar_producto(COCA)
        self.assertEqual(self.a.cambiar_precios(10, "Bebidas"), 1)
        self.assertEqual(self.a.producto(COCA["codigo"])["precio"], 1320)
        self.assertEqual(self.a.producto(codigo)["precio"], 100)
        self.a.dar_de_baja(codigo)
        self.assertEqual([p["codigo"] for p in self.a.lista_productos()], [COCA["codigo"]])
        self.assertEqual(len(self.a.lista_productos(incluir_bajas=True)), 2)
        self.assertEqual(len(self.a.lista_productos("coca")), 1)
        self.assertEqual(len(self.a.lista_productos("panaderia")), 0)  # dado de baja
        self.reabrir()
        self.assertEqual(self.a.producto(codigo)["activo"], "NO")

    def test_registro_tolera_linea_cortada(self):
        self.a.guardar_producto(COCA)
        registro = next((self.carpeta / "registro").glob("*.jsonl"))
        with open(registro, "a", encoding="utf-8") as f:
            f.write('{"seq": 99, "fecha": "2026')  # corte de luz a mitad de línea
        self.reabrir(guardar=False)
        self.assertIsNotNone(self.a.producto(COCA["codigo"]))
        self.assertEqual(json.loads(registro.read_text().splitlines()[0])["tipo"], "producto")

    def test_exportar_sin_proteccion(self):
        self.a.guardar_producto(COCA)
        self.a.abrir_caja(0)
        venta = self.a.registrar_venta([item(COCA["codigo"], "Coca", 1, 1200)], "Efectivo")
        dia = venta["fecha"][:10]
        ruta = self.a.exportar(dia, dia)
        libro = load_workbook(ruta)
        self.assertFalse(libro["Ventas"].protection.sheet)
        self.assertEqual(libro["Ventas"].max_row, 2)


if __name__ == "__main__":
    unittest.main()
